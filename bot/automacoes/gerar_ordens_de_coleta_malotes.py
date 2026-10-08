from __future__ import annotations

import json
import re
from datetime import datetime
from pathlib import Path
from typing import Callable

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from playwright.sync_api import (
    Locator,
    Page,
    TimeoutError as PlaywrightTimeoutError,
    sync_playwright,
)

from servicos.fraction import (
    fechar_recurso_playwright_seguro,
    recuperar_pagina_fraction,
)
from utils.config import (
    ARQUIVO_BASE_CNPJS,
    ARQUIVO_EMAILS_UNIDADES,
    ARQUIVO_ENV,
    FRACTION_PASSWORD,
    FRACTION_USER,
    RESULTADOS_ORDENS,
    URL_FRACTION,
)


PASTA_RESULTADOS = RESULTADOS_ORDENS
ARQUIVO_EMAILS = ARQUIVO_EMAILS_UNIDADES
URL = URL_FRACTION

TIMEOUT = 30_000

MAX_TENTATIVAS = 5

CNPJ_DESTINATARIO = "42591651000143"

# Valores definidos durante os testes do formulário.
CONTA_CORRENTE = "0153080"
OBSERVACAO = "bag"
CONTEUDO = "bag - malote"
VOLUMES = "1"
PESO = "1,00"
VALOR_COLETA = "13,20"
NUMERO_NOTA = "dec"
SERIE = "0"
VALOR_NOTA = "100,00"
MODALIDADE = "CORPORATE"

def carregar_emails_unidades() -> dict[str, str]:
    if not ARQUIVO_EMAILS.exists():
        raise FileNotFoundError(
            f"Base de e-mails não encontrada: {ARQUIVO_EMAILS}"
        )

    with ARQUIVO_EMAILS.open(
        "r",
        encoding="utf-8",
    ) as arquivo:
        dados = json.load(arquivo)

    if not isinstance(dados, dict):
        raise ValueError(
            "emails_unidades.json deve possuir o formato "
            '{"UNIDADE": ["email@dominio.com.br"]}.'
        )

    emails: dict[str, str] = {}

    for unidade, valores in dados.items():
        unidade = texto_limpo(unidade)

        if not unidade:
            continue

        if isinstance(valores, list):
            lista_emails = [
                texto_limpo(email)
                for email in valores
                if texto_limpo(email)
            ]
        elif isinstance(valores, str):
            lista_emails = [
                texto_limpo(valores)
            ] if texto_limpo(valores) else []
        else:
            raise ValueError(
                f"Formato de e-mail inválido para a unidade {unidade}."
            )

        emails[unidade] = "; ".join(lista_emails)

    return emails

def _criar_planilha_resumo_ordens(
    detalhes: list[dict],
    situacao_alvo: str,
    nome_arquivo: str,
    titulo_aba: str,
    *,
    pronta_reimportacao: bool,
) -> Path | None:
    itens = [
        item
        for item in detalhes
        if item["SITUAÇÃO"] == situacao_alvo
    ]

    if not itens:
        return None

    df_itens = pd.DataFrame(
        itens
    )

    resumo = (
        df_itens
        .groupby(
            "SIGLA",
            as_index=False,
        )
        .size()
        .rename(
            columns={
                "SIGLA": "SIGLA DO RESTAURANTE",
                "size": "NUMERO DE ORDENS",
            }
        )
    )

    PASTA_RESULTADOS.mkdir(
        parents=True,
        exist_ok=True,
    )

    agora = datetime.now()

    caminho = (
        PASTA_RESULTADOS
        / agora.strftime(
            nome_arquivo
        )
    )

    wb = Workbook()
    ws = wb.active
    ws.title = titulo_aba

    if pronta_reimportacao:
        cabecalhos = [
            "SIGLA DO RESTAURANTE",
            "NUMERO DE ORDENS",
        ]
    else:
        cabecalhos = [
            "SIGLA DO RESTAURANTE",
            "NUMERO DE ORDENS",
            "SITUAÇÃO",
            "AÇÃO RECOMENDADA",
            "MENSAGEM",
        ]

    ws.append(
        cabecalhos
    )

    mensagens_por_sigla = (
        df_itens
        .groupby("SIGLA")["MENSAGEM"]
        .apply(
            lambda valores: " | ".join(
                dict.fromkeys(
                    texto_limpo(valor)
                    for valor in valores
                    if texto_limpo(valor)
                )
            )
        )
        .to_dict()
    )

    for _, linha in resumo.iterrows():
        sigla = linha[
            "SIGLA DO RESTAURANTE"
        ]

        quantidade = int(
            linha["NUMERO DE ORDENS"]
        )

        if pronta_reimportacao:
            ws.append(
                [
                    sigla,
                    quantidade,
                ]
            )
        else:
            ws.append(
                [
                    sigla,
                    quantidade,
                    "STATUS INDETERMINADO",
                    (
                        "Conferir no Fraction se a coleta foi criada "
                        "antes de reimportar."
                    ),
                    mensagens_por_sigla.get(
                        sigla,
                        "",
                    ),
                ]
            )

    azul_cabecalho = PatternFill(
        fill_type="solid",
        fgColor="8EDDE1",
    )

    if pronta_reimportacao:
        cor_linha = "FFC7CE"
        cor_fonte = "9C0006"
    else:
        cor_linha = "FFF2CC"
        cor_fonte = "7F6000"

    preenchimento = PatternFill(
        fill_type="solid",
        fgColor=cor_linha,
    )

    fonte_linha = Font(
        color=cor_fonte,
    )

    borda_fina = Border(
        left=Side(
            style="thin",
            color="000000",
        ),
        right=Side(
            style="thin",
            color="000000",
        ),
        top=Side(
            style="thin",
            color="000000",
        ),
        bottom=Side(
            style="thin",
            color="000000",
        ),
    )

    for celula in ws[1]:
        celula.fill = azul_cabecalho
        celula.font = Font(
            bold=True,
        )
        celula.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )
        celula.border = borda_fina

    for numero_linha in range(
        2,
        ws.max_row + 1,
    ):
        for numero_coluna in range(
            1,
            ws.max_column + 1,
        ):
            celula = ws.cell(
                numero_linha,
                numero_coluna,
            )

            celula.fill = preenchimento
            celula.font = fonte_linha
            celula.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )
            celula.border = borda_fina

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 22

    if not pronta_reimportacao:
        ws.column_dimensions["C"].width = 25
        ws.column_dimensions["D"].width = 55
        ws.column_dimensions["E"].width = 75

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = (
        f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    )

    wb.save(
        caminho
    )

    return caminho


def criar_planilha_erros_reimportar(
    detalhes: list[dict],
) -> Path | None:
    """
    Somente falhas ocorridas antes do clique em Gerar Coleta.

    Este arquivo é seguro para importar novamente no mesmo fluxo.
    """
    return _criar_planilha_resumo_ordens(
        detalhes,
        "FALHA",
        "ordens_erros_reimportar_%d-%m-%Y_%H-%M-%S.xlsx",
        "Erros para Reimportar",
        pronta_reimportacao=True,
    )


def criar_planilha_verificar(
    detalhes: list[dict],
) -> Path | None:
    """
    Casos em que a falha aconteceu depois do clique em Gerar Coleta.

    A coleta pode ter sido criada no Fraction. Por segurança, estes itens
    NÃO entram na planilha pronta para reimportação.
    """
    return _criar_planilha_resumo_ordens(
        detalhes,
        "STATUS INDETERMINADO",
        "ordens_verificar_%d-%m-%Y_%H-%M-%S.xlsx",
        "Verificar no Fraction",
        pronta_reimportacao=False,
    )


def normalizar_coluna(valor: object) -> str:
    return " ".join(str(valor).strip().upper().split())


def normalizar_sigla(valor: object) -> str:
    return " ".join(str(valor).strip().upper().split())


def texto_limpo(valor: object) -> str:
    if valor is None:
        return ""

    texto = str(valor).strip()

    if texto.lower() == "nan":
        return ""

    return texto


def somente_digitos(valor: object) -> str:
    return "".join(
        caractere
        for caractere in texto_limpo(valor)
        if caractere.isdigit()
    )


def ler_excel_normalizado(caminho_ou_arquivo) -> pd.DataFrame:
    df = pd.read_excel(
        caminho_ou_arquivo,
        dtype=str,
    ).fillna("")

    df.columns = [
        normalizar_coluna(coluna)
        for coluna in df.columns
    ]

    return df


def validar_arquivos_fixos() -> list[str]:
    erros: list[str] = []

    if not ARQUIVO_ENV.exists():
        erros.append(
            "Crie o arquivo .env na raiz do projeto."
        )
    else:
        usuario = FRACTION_USER
        senha = FRACTION_PASSWORD
        url_fraction = URL_FRACTION

        if not usuario:
            erros.append(
                "FRACTION_USER não foi informado no .env."
            )

        if not senha:
            erros.append(
                "FRACTION_PASSWORD não foi informado no .env."
            )

        if not url_fraction:
            erros.append(
                "URL_FRACTION não foi informada no .env."
            )

    if not ARQUIVO_BASE_CNPJS.exists():
        erros.append(
            "Crie dados/base_cnpjs.json."
        )
    else:
        try:
            carregar_base_cnpjs()
        except Exception as erro:
            erros.append(
                f"Não foi possível ler base_cnpjs.json: {erro}"
            )

    if not ARQUIVO_EMAILS.exists():
        erros.append(
            "Crie dados/emails_unidades.json."
        )
    else:
        try:
            carregar_emails_unidades()
        except Exception as erro:
            erros.append(
                f"Não foi possível ler emails_unidades.json: {erro}"
            )

    return erros

def validar_planilha_pedidos(df_original: pd.DataFrame) -> list[str]:
    erros: list[str] = []

    df = df_original.copy()
    df.columns = [
        normalizar_coluna(coluna)
        for coluna in df.columns
    ]

    obrigatorias = {
        "SIGLA DO RESTAURANTE",
        "NUMERO DE ORDENS",
    }

    faltantes = obrigatorias - set(df.columns)

    if faltantes:
        erros.append(
            "Colunas ausentes: "
            + ", ".join(sorted(faltantes))
        )
        return erros

    if df.empty:
        erros.append("A base de pedidos está vazia.")
        return erros

    for indice, linha in df.iterrows():
        numero_linha = indice + 2
        sigla = normalizar_sigla(
            linha["SIGLA DO RESTAURANTE"]
        )
        quantidade_texto = texto_limpo(
            linha["NUMERO DE ORDENS"]
        )

        if not sigla:
            erros.append(
                f"Linha {numero_linha}: sigla vazia."
            )

        try:
            quantidade = int(float(quantidade_texto))
        except (TypeError, ValueError):
            erros.append(
                f"Linha {numero_linha}: NUMERO DE ORDENS inválido."
            )
            continue

        if quantidade <= 0:
            erros.append(
                f"Linha {numero_linha}: NUMERO DE ORDENS "
                "deve ser maior que zero."
            )

    return erros


def carregar_login() -> tuple[str, str]:
    return FRACTION_USER, FRACTION_PASSWORD


def carregar_base_cnpjs() -> dict[str, str]:
    if not ARQUIVO_BASE_CNPJS.exists():
        raise FileNotFoundError(
            f"Arquivo não encontrado: {ARQUIVO_BASE_CNPJS}"
        )

    with ARQUIVO_BASE_CNPJS.open(
        "r",
        encoding="utf-8",
    ) as arquivo:
        dados = json.load(arquivo)

    if not isinstance(dados, dict):
        raise ValueError(
            "base_cnpjs.json deve possuir o formato "
            '{"SIGLA": "CNPJ"}.'
        )

    base: dict[str, str] = {}

    for sigla, cnpj in dados.items():
        sigla = normalizar_sigla(sigla)
        cnpj = somente_digitos(cnpj)

        if not sigla:
            continue

        if len(cnpj) != 14:
            raise ValueError(
                f"CNPJ inválido para a sigla {sigla}: {cnpj}"
            )

        if sigla in base:
            raise ValueError(
                f"A sigla {sigla} aparece mais de uma vez "
                "em base_cnpjs.json."
            )

        base[sigla] = cnpj

    return base

def preparar_execucoes(
    df_original: pd.DataFrame,
) -> tuple[list[dict], list[str]]:
    df = df_original.copy()
    df.columns = [
        normalizar_coluna(coluna)
        for coluna in df.columns
    ]

    base_cnpjs = carregar_base_cnpjs()

    execucoes: list[dict] = []
    alertas: list[str] = []

    for indice, linha in df.iterrows():
        sigla = normalizar_sigla(
            linha["SIGLA DO RESTAURANTE"]
        )
        quantidade = int(
            float(texto_limpo(linha["NUMERO DE ORDENS"]))
        )

        cnpj = base_cnpjs.get(sigla)

        if cnpj is None:
            alertas.append(
                f"Linha {indice + 2}: sigla {sigla} não encontrada "
                "em base_cnpjs.json."
            )
            continue

        for sequencia in range(1, quantidade + 1):
            execucoes.append(
                {
                    "sigla": sigla,
                    "sequencia": sequencia,
                    "total_sigla": quantidade,
                    "cnpj_remetente": cnpj,
                }
            )

    return execucoes, alertas


def digitar(campo: Locator, valor: str) -> None:
    campo.wait_for(
        state="visible",
        timeout=TIMEOUT,
    )
    campo.scroll_into_view_if_needed()
    campo.click()
    campo.press("Control+A")
    campo.press("Backspace")
    campo.type(
        texto_limpo(valor),
        delay=50,
    )


def colar_sem_tab(
    page: Page,
    campo: Locator,
    valor: str,
) -> None:
    campo.wait_for(
        state="visible",
        timeout=TIMEOUT,
    )
    campo.scroll_into_view_if_needed()
    campo.click()
    campo.press("Control+A")
    campo.press("Backspace")

    page.evaluate(
        "valor => navigator.clipboard.writeText(valor)",
        texto_limpo(valor),
    )

    campo.press("Control+V")
    page.wait_for_timeout(400)


def realizar_login(
    page: Page,
    usuario: str,
    senha: str,
    log: Callable[[str], None],
) -> None:
    log("Abrindo Jadlog...")

    page.goto(
        URL,
        wait_until="domcontentloaded",
        timeout=60_000,
    )

    try:
        campo_usuario = page.get_by_role(
            "textbox",
            name="Usuário",
            exact=True,
        )

        campo_usuario.wait_for(
            state="visible",
            timeout=7_000,
        )

        log("Realizando login...")

        campo_usuario.fill(usuario)

        page.get_by_role(
            "textbox",
            name="Senha",
            exact=True,
        ).fill(senha)

        page.get_by_role("button").click()

    except PlaywrightTimeoutError:
        log(
            "Tela de login não apareceu. "
            "A sessão pode já estar autenticada."
        )

    page.get_by_role(
        "link",
        name="Operacional",
    ).wait_for(
        state="visible",
        timeout=60_000,
    )

    log("Login concluído.")


def abrir_menu_com_submenu(
    menu: Locator,
    submenu: Locator,
    nome_menu: str,
) -> None:
    menu.wait_for(
        state="visible",
        timeout=TIMEOUT,
    )

    for tentativa in range(1, 4):
        menu.click()

        try:
            submenu.wait_for(
                state="visible",
                timeout=3_000,
            )
            return

        except PlaywrightTimeoutError:
            if tentativa == 3:
                raise RuntimeError(
                    f"Não foi possível abrir o menu {nome_menu}."
                )

def abrir_solicitacao_coleta(
    page: Page,
    log: Callable[[str], None],
) -> None:
    log("Abrindo Operacional...")

    menu_operacional = page.get_by_role(
        "link",
        name="Operacional",
    )

    submenu_ordem = page.get_by_role(
        "link",
        name="Ordem de coleta",
    )

    abrir_menu_com_submenu(
        menu_operacional,
        submenu_ordem,
        "Operacional",
    )

    log("Abrindo Ordem de coleta...")

    submenu_solicitacao = page.get_by_role(
        "link",
        name="Solicitação de Coleta",
    )

    abrir_menu_com_submenu(
        submenu_ordem,
        submenu_solicitacao,
        "Ordem de coleta",
    )

    log("Abrindo Solicitação de Coleta...")

    submenu_solicitacao.click()

    page.get_by_role(
        "textbox",
        name="Conta Corrente",
        exact=True,
    ).wait_for(
        state="visible",
        timeout=TIMEOUT,
    )

    log("Tela de Solicitação de Coleta aberta.")


def localizar_remetente(page: Page) -> Locator:
    return (
        page.get_by_role(
            "cell",
            name="Remetente CNPJ/CPF: Insc.",
        )
        .get_by_label("CNPJ/CPF:")
    )


def localizar_destinatario(page: Page) -> Locator:
    return (
        page.get_by_role(
            "cell",
            name="Destinatário CNPJ/CPF: Insc.",
        )
        .get_by_label("CNPJ/CPF:")
    )


def selecionar_modalidade_corporate(page: Page) -> None:
    modalidade = page.locator(
        '[id="form_emissao:modalidadeSelect"]'
    )

    modalidade.wait_for(
        state="visible",
        timeout=TIMEOUT,
    )

    caixa = modalidade.bounding_box()

    if caixa is None:
        raise RuntimeError(
            "Não foi possível localizar o campo Modalidade."
        )

    page.mouse.click(
        caixa["x"] + caixa["width"] - 10,
        caixa["y"] + caixa["height"] / 2,
    )

    page.wait_for_timeout(600)

    lista = page.locator(
        '[id="form_emissao:modalidadeSelect_items"]:visible'
    )

    lista.wait_for(
        state="visible",
        timeout=10_000,
    )

    lista.get_by_role(
        "option",
        name=MODALIDADE,
        exact=True,
    ).click()

    page.wait_for_timeout(700)


def preencher_formulario(
    page: Page,
    cnpj_remetente: str,
) -> None:
    digitar(
        page.get_by_role(
            "textbox",
            name="Conta Corrente",
            exact=True,
        ),
        CONTA_CORRENTE,
    )

    digitar(
        page.locator(
            '[id="form_emissao:observacaoArea"]'
        ),
        OBSERVACAO,
    )

    digitar(
        page.locator(
            '[id="form_emissao:conteudoArea"]'
        ),
        CONTEUDO,
    )

    # Seletores validados durante os testes.
    digitar(
        page.get_by_role(
            "textbox",
            name="ALT+7",
            exact=True,
        ),
        PESO,
    )

    digitar(
        page.get_by_role(
            "textbox",
            name="ALT+8",
            exact=True,
        ),
        VALOR_COLETA,
    )

    digitar(
        page.get_by_role(
            "textbox",
            name="ALT+9",
            exact=True,
        ),
        NUMERO_NOTA,
    )

    # Volume e série já aparecem com 1 e 0 na página.
    # Os valores são reforçados diretamente quando os IDs existem.
    campo_volume = page.locator(
        '[id="form_emissao:quantidadeVolume_input"]'
    )

    if campo_volume.count() > 0:
        digitar(campo_volume, VOLUMES)

    campo_serie = page.locator(
        '[id="form_emissao:bonfs:0:nota_serie"]'
    )

    if campo_serie.count() > 0:
        digitar(campo_serie, SERIE)

    digitar(
        page.locator(
            '[id="form_emissao:bonfs:0:nota_valor"]'
        ),
        VALOR_NOTA,
    )

    # Os dois CNPJs são colados antes de sair dos campos.
    colar_sem_tab(
        page,
        localizar_remetente(page),
        cnpj_remetente,
    )

    colar_sem_tab(
        page,
        localizar_destinatario(page),
        CNPJ_DESTINATARIO,
    )

    # Modalidade por último. O clique fora dispara o carregamento
    # dos dados de remetente e destinatário.
    selecionar_modalidade_corporate(page)

    page.wait_for_timeout(4_000)

    modalidade_exibida = page.locator(
        '[id="form_emissao:modalidadeSelect_label"]'
    ).inner_text()

    if modalidade_exibida.strip().upper() != MODALIDADE:
        raise RuntimeError(
            f"Modalidade inesperada: {modalidade_exibida}"
        )

def capturar_unidade_coletora(page: Page) -> str:
    texto_unidade = page.get_by_text(
        "Unidade coletora",
        exact=True,
    ).first

    texto_unidade.wait_for(
        state="visible",
        timeout=TIMEOUT,
    )

    celula = texto_unidade.locator(
        "xpath=ancestor::td[1]"
    )

    texto_completo = celula.inner_text().strip()

    linhas = [
        linha.strip()
        for linha in texto_completo.splitlines()
        if linha.strip()
        and linha.strip().lower() != "unidade coletora"
    ]

    if linhas:
        unidade = " ".join(linhas)

        return re.sub(
            r"\s+",
            " ",
            unidade,
        ).strip()

    # Alternativa para o caso de o nome estar
    # em uma célula anterior.
    celula_anterior = texto_unidade.locator(
        "xpath=ancestor::td[1]/preceding-sibling::td[1]"
    )

    if celula_anterior.count() > 0:
        unidade = celula_anterior.inner_text().strip()

        if unidade:
            return re.sub(
                r"\s+",
                " ",
                unidade,
            ).strip()

    raise RuntimeError(
        "A Unidade coletora apareceu, "
        "mas o nome da unidade não foi capturado."
    )

def gerar_coleta_e_capturar_ordem(page: Page) -> str:
    page.get_by_role(
        "button",
        name="Gerar Coleta",
        exact=True,
    ).click()

    mensagem = page.get_by_text(
        re.compile(
            r"Gerado\s+o\s+número\s+de\s+coleta",
            re.IGNORECASE,
        )
    ).last

    mensagem.wait_for(
        state="visible",
        timeout=30_000,
    )

    texto = mensagem.inner_text()

    resultado = re.search(
        r"Gerado\s+o\s+número\s+de\s+coleta\s+(\d+)",
        texto,
        flags=re.IGNORECASE,
    )

    if not resultado:
        # Fallback no contêiner branco mapeado pelo usuário.
        caixa = (
            page.locator("div")
            .filter(
                has_text=re.compile(
                    r"Gerado\s+o\s+número\s+de\s+coleta",
                    re.IGNORECASE,
                )
            )
            .nth(2)
        )

        caixa.wait_for(
            state="visible",
            timeout=10_000,
        )

        texto = caixa.inner_text()

        resultado = re.search(
            r"Gerado\s+o\s+número\s+de\s+coleta\s+(\d+)",
            texto,
            flags=re.IGNORECASE,
        )

    if not resultado:
        raise RuntimeError(
            "A coleta foi enviada, mas o número não pôde ser "
            f"extraído da mensagem: {texto}"
        )

    return resultado.group(1)


def criar_planilha_resultados(
    resultados: list[dict],
) -> Path:
    PASTA_RESULTADOS.mkdir(
        parents=True,
        exist_ok=True,
    )

    agora = datetime.now()
    nome = agora.strftime(
        "novas_ordens_%d-%m-%Y_%H-%M-%S.xlsx"
    )
    caminho = PASTA_RESULTADOS / nome

    wb = Workbook()
    ws = wb.active
    ws.title = "Novas Ordens"

    cabecalhos = [
        "RE",
        "SIGLA",
        "TIPO",
        "CTE",
        "VINCULAR/ ACERTO",
        "ORDEM",
        "SITUAÇÃO",
        "DT DE FINALIZAÇÃO",
        "DIAS FALTANTES",
        "SITUAÇÃO DE COLETA",
        "UNIDADE",
        "E-MAIL",
        "MÊS",
        "ANO",
        "SITUAÇÃO DO PEDIDO",
    ]

    ws.append(cabecalhos)

    for item in resultados:
        ws.append(
            [
                item["RE"],
                item["SIGLA"],
                "",                   # TIPO
                "",                   # CTE
                "",                   # VINCULAR/ ACERTO
                item["ORDEM"],
                "AUTORIZADO",
                "",                   # DT DE FINALIZAÇÃO
                "",                   # DIAS FALTANTES
                "",                   # SITUAÇÃO DE COLETA
                item["UNIDADE"],
                item["EMAIL"],        # E-MAIL
                "",                   # MÊS
                "",                   # ANO
                "",                   # SITUAÇÃO DO PEDIDO
            ]
        )

    azul_cabecalho = PatternFill(
        fill_type="solid",
        fgColor="8EDDE1",
    )
    verde_autorizado = PatternFill(
        fill_type="solid",
        fgColor="C6EFCE",
    )
    fonte_autorizado = Font(
        color="006100",
    )
    borda_fina = Border(
        left=Side(style="thin", color="000000"),
        right=Side(style="thin", color="000000"),
        top=Side(style="thin", color="000000"),
        bottom=Side(style="thin", color="000000"),
    )

    for celula in ws[1]:
        celula.fill = azul_cabecalho
        celula.font = Font(bold=True)
        celula.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )
        celula.border = borda_fina

    for linha in range(2, ws.max_row + 1):
        for coluna in range(1, 16):
            celula = ws.cell(linha, coluna)
            celula.border = borda_fina
            celula.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        situacao = ws.cell(linha, 7)
        situacao.fill = verde_autorizado
        situacao.font = fonte_autorizado

    larguras = {
        "A": 14,   # RE
        "B": 18,   # SIGLA
        "C": 14,   # TIPO
        "D": 14,   # CTE
        "E": 24,   # VINCULAR/ ACERTO
        "F": 18,   # ORDEM
        "G": 18,   # SITUAÇÃO
        "H": 22,   # DT DE FINALIZAÇÃO
        "I": 18,   # DIAS FALTANTES
        "J": 24,   # SITUAÇÃO DE COLETA
        "K": 38,   # UNIDADE
        "L": 28,   # E-MAIL
        "M": 14,   # MÊS
        "N": 12,   # ANO
        "O": 24,   # SITUAÇÃO DO PEDIDO
    }

    for coluna, largura in larguras.items():
        ws.column_dimensions[coluna].width = largura

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:O{ws.max_row}"

    wb.save(caminho)

    return caminho


def reiniciar_sessao_coleta(
    page: Page,
    usuario: str,
    senha: str,
    log: Callable[[str], None],
) -> None:
    """
    Limpa o estado da Solicitação de Coleta antes de uma nova tentativa.

    O reset principal é feito pelo logout explícito do Fraction:
    Sair -> tela de login -> novo login.

    Isso evita reaproveitar uma tela de Solicitação de Coleta que manteve
    campos preenchidos ou estado interno após uma falha.
    """
    log(
        "Encerrando a sessão atual do Fraction antes da nova tentativa..."
    )

    sair = page.get_by_role(
        "link",
        name="Sair",
    )

    sair.wait_for(
        state="visible",
        timeout=TIMEOUT,
    )

    sair.click(
        timeout=TIMEOUT,
    )

    campo_usuario = page.get_by_role(
        "textbox",
        name="Usuário",
        exact=True,
    )

    campo_usuario.wait_for(
        state="visible",
        timeout=TIMEOUT,
    )

    log(
        "Sessão encerrada. Realizando um novo login..."
    )

    realizar_login(
        page,
        usuario,
        senha,
        log,
    )

    log(
        "Nova sessão do Fraction iniciada."
    )


def recuperar_home(
    page: Page,
    usuario: str,
    senha: str,
    log: Callable[[str], None],
) -> None:
    def validar_operacional() -> None:
        page.get_by_role(
            "link",
            name="Operacional",
        ).wait_for(
            state="visible",
            timeout=60_000,
        )

    recuperar_pagina_fraction(
        page,
        log,
        validar_callback=validar_operacional,
        login_callback=lambda: realizar_login(
            page,
            usuario,
            senha,
            log,
        ),
        nome_tela="Operacional do Fraction",
    )


def executar_automacao(
    execucoes: list[dict],
    headless: bool,
    continuar_em_erro: bool,
    log: Callable[[str], None],
) -> dict:
    erros_fixos = validar_arquivos_fixos()

    emails_unidades = carregar_emails_unidades()

    if erros_fixos:
        raise ValueError("\n".join(erros_fixos))

    usuario, senha = carregar_login()

    autorizadas: list[dict] = []
    detalhes: list[dict] = []
    sucessos = 0
    falhas = 0
    indeterminados = 0

    with sync_playwright() as p:
        navegador = None
        contexto = None
        page = None

        def fechar_sessao() -> None:
            nonlocal navegador, contexto, page

            fechar_recurso_playwright_seguro(
                contexto,
                log,
            )
            fechar_recurso_playwright_seguro(
                navegador,
                log,
            )

            navegador = None
            contexto = None
            page = None

        def abrir_sessao() -> None:
            nonlocal navegador, contexto, page

            fechar_sessao()

            navegador = p.chromium.launch(
                headless=headless,
                slow_mo=120 if not headless else 0,
            )

            contexto = navegador.new_context(
                permissions=[
                    "clipboard-read",
                    "clipboard-write",
                ],
            )

            page = contexto.new_page()
            page.set_default_timeout(TIMEOUT)

            realizar_login(
                page,
                usuario,
                senha,
                log,
            )

        try:
            try:
                abrir_sessao()
            except Exception as erro_login:
                mensagem = (
                    "Fraction indisponível antes de iniciar as ordens: "
                    f"{type(erro_login).__name__}: {erro_login}"
                )

                log(mensagem)

                for item in execucoes:
                    detalhes.append(
                        {
                            "SIGLA": item["sigla"],
                            "ORDEM": "",
                            "SITUAÇÃO": "FALHA",
                            "MENSAGEM": mensagem,
                        }
                    )

                falhas = len(execucoes)

                arquivo_erros = criar_planilha_erros_reimportar(
                    detalhes
                )

                return {
                    "total": len(execucoes),
                    "sucessos": 0,
                    "falhas": falhas,
                    "indeterminados": 0,
                    "detalhes": detalhes,
                    "arquivo_resultado": None,
                    "arquivo_erros": (
                        str(arquivo_erros)
                        if arquivo_erros
                        else None
                    ),
                    "arquivo_verificar": None,
                }

            total = len(execucoes)
            interromper_restantes = False

            for posicao, item in enumerate(
                execucoes,
                start=1,
            ):
                sigla = item["sigla"]

                if interromper_restantes:
                    restantes = execucoes[
                        posicao - 1:
                    ]

                    mensagem = (
                        "Fraction indisponível após tentativas de "
                        "recuperação. Item não executado."
                    )

                    for restante in restantes:
                        detalhes.append(
                            {
                                "SIGLA": restante["sigla"],
                                "ORDEM": "",
                                "SITUAÇÃO": "FALHA",
                                "MENSAGEM": mensagem,
                            }
                        )

                    falhas += len(restantes)
                    log(
                        f"{len(restantes)} ordem(ns) restante(s) "
                        "foram enviadas para a planilha de erros."
                    )
                    break

                log("")
                log(
                    f"[{posicao}/{total}] {sigla} "
                    f"({item['sequencia']}/{item['total_sigla']})"
                )

                sucesso_item = False
                status_indeterminado = False
                ultimo_erro = None
                ultima_mensagem = ""

                for tentativa in range(
                    1,
                    MAX_TENTATIVAS + 1,
                ):
                    geracao_iniciada = False

                    log(
                        f"Tentativa {tentativa}/"
                        f"{MAX_TENTATIVAS}"
                    )

                    try:
                        # -----------------------------------------
                        # CAMINHO COMPLETO
                        # -----------------------------------------

                        abrir_solicitacao_coleta(
                            page,
                            log,
                        )

                        log(
                            f"CNPJ do remetente: "
                            f"{item['cnpj_remetente']}"
                        )

                        preencher_formulario(
                            page,
                            item["cnpj_remetente"],
                        )

                        # -----------------------------------------
                        # UNIDADE COLETORA
                        # -----------------------------------------

                        log(
                            "Capturando Unidade coletora..."
                        )

                        unidade_coletora = (
                            capturar_unidade_coletora(
                                page
                            )
                        )

                        log(
                            f"Unidade coletora: "
                            f"{unidade_coletora}"
                        )

                        # -----------------------------------------
                        # E-MAIL DA UNIDADE
                        # -----------------------------------------

                        email_unidade = (
                            emails_unidades.get(
                                unidade_coletora,
                                "",
                            )
                        )

                        if email_unidade:
                            log(
                                f"E-mail da unidade: "
                                f"{email_unidade}"
                            )
                        else:
                            log(
                                "AVISO: E-mail não encontrado "
                                f"para {unidade_coletora}"
                            )

                        # -----------------------------------------
                        # GERAÇÃO DA ORDEM
                        # -----------------------------------------
                        #
                        # A partir daqui NÃO fazemos nova
                        # tentativa automática caso aconteça erro,
                        # pois o clique pode ter gerado a coleta.
                        # -----------------------------------------

                        geracao_iniciada = True

                        numero_ordem = (
                            gerar_coleta_e_capturar_ordem(
                                page
                            )
                        )

                        data_geracao = (
                            datetime.now().strftime(
                                "%d/%m/%Y"
                            )
                        )

                        autorizadas.append(
                            {
                                "RE": data_geracao,
                                "SIGLA": sigla,
                                "ORDEM": numero_ordem,
                                "UNIDADE": (
                                    unidade_coletora
                                ),
                                "EMAIL": email_unidade,
                            }
                        )

                        detalhes.append(
                            {
                                "SIGLA": sigla,
                                "ORDEM": numero_ordem,
                                "UNIDADE": (
                                    unidade_coletora
                                ),
                                "SITUAÇÃO": "AUTORIZADO",
                                "MENSAGEM": (
                                    "Coleta gerada "
                                    "com sucesso."
                                ),
                            }
                        )

                        sucessos += 1
                        sucesso_item = True

                        log(
                            f"Ordem gerada: "
                            f"{numero_ordem}"
                        )

                        break

                    except Exception as erro:
                        ultimo_erro = erro

                        ultima_mensagem = (
                            f"{type(erro).__name__}: "
                            f"{erro}"
                        )

                        log(
                            f"Falha na tentativa "
                            f"{tentativa}/"
                            f"{MAX_TENTATIVAS}: "
                            f"{ultima_mensagem}"
                        )

                        # -----------------------------------------
                        # NÃO RETENTA SE JÁ INICIOU A GERAÇÃO
                        # -----------------------------------------

                        if geracao_iniciada:
                            status_indeterminado = True

                            log(
                                "A falha ocorreu após o clique em "
                                "Gerar Coleta."
                            )

                            log(
                                "A coleta pode ter sido criada no Fraction. "
                                "Este item não será repetido automaticamente."
                            )

                            log(
                                "O item será enviado para a planilha "
                                "de verificação, separada da planilha "
                                "pronta para reimportação."
                            )

                            break

                        # -----------------------------------------
                        # AINDA TEM TENTATIVA?
                        # -----------------------------------------

                        if (
                            tentativa
                            < MAX_TENTATIVAS
                        ):
                            log(
                                "Preparando nova tentativa..."
                            )

                            try:
                                reiniciar_sessao_coleta(
                                    page,
                                    usuario,
                                    senha,
                                    log,
                                )

                            except Exception as erro_logout:
                                log(
                                    "Não foi possível limpar a sessão usando "
                                    "o botão Sair. Fechando o Chromium e "
                                    "abrindo uma sessão nova."
                                )

                                try:
                                    abrir_sessao()

                                except Exception as erro_restart:
                                    ultima_mensagem = (
                                        "Falha ao restabelecer o Fraction: "
                                        f"{type(erro_restart).__name__}: "
                                        f"{erro_restart}"
                                    )

                                    log(
                                        ultima_mensagem
                                    )

                                    ultimo_erro = erro_restart
                                    interromper_restantes = True
                                    break

                            continue

                        log(
                            "Número máximo de tentativas "
                            "atingido."
                        )

                # =============================================
                # SÓ REGISTRA FALHA DEPOIS DAS TENTATIVAS
                # =============================================

                if not sucesso_item:
                    if status_indeterminado:
                        indeterminados += 1

                        detalhes.append(
                            {
                                "SIGLA": sigla,
                                "ORDEM": "",
                                "SITUAÇÃO": "STATUS INDETERMINADO",
                                "MENSAGEM": (
                                    ultima_mensagem
                                ),
                            }
                        )

                        log(
                            f"Ordem da sigla {sigla} marcada como "
                            "STATUS INDETERMINADO."
                        )

                    else:
                        falhas += 1

                        detalhes.append(
                            {
                                "SIGLA": sigla,
                                "ORDEM": "",
                                "SITUAÇÃO": "FALHA",
                                "MENSAGEM": (
                                    ultima_mensagem
                                ),
                            }
                        )

                        log(
                            f"Ordem da sigla {sigla} "
                            "marcada como FALHA."
                        )

                    if not continuar_em_erro:
                        raise RuntimeError(
                            ultima_mensagem
                        ) from ultimo_erro

                    # Limpa totalmente a sessão antes
                    # de seguir para a próxima ordem.
                    try:
                        reiniciar_sessao_coleta(
                            page,
                            usuario,
                            senha,
                            log,
                        )

                    except Exception:
                        log(
                            "Não foi possível encerrar e refazer a sessão "
                            "após a falha final. Reiniciando o Chromium "
                            "para a próxima ordem."
                        )

                        try:
                            abrir_sessao()
                            interromper_restantes = False
                        except Exception as erro_restart:
                            log(
                                "Não foi possível restabelecer o Fraction: "
                                f"{type(erro_restart).__name__}: {erro_restart}"
                            )
                            interromper_restantes = True

        finally:
            fechar_sessao()

    arquivo_resultado = None
    arquivo_erros = None
    arquivo_verificar = None

    if autorizadas:
        arquivo_resultado = criar_planilha_resultados(
            autorizadas
        )

    arquivo_erros = criar_planilha_erros_reimportar(
        detalhes
    )

    arquivo_verificar = criar_planilha_verificar(
        detalhes
    )

    return {
        "total": len(execucoes),
        "sucessos": sucessos,
        "falhas": falhas,
        "indeterminados": indeterminados,
        "detalhes": detalhes,

        "arquivo_resultado": (
            str(arquivo_resultado)
            if arquivo_resultado
            else None
        ),

        "arquivo_erros": (
            str(arquivo_erros)
            if arquivo_erros
            else None
        ),

        "arquivo_verificar": (
            str(arquivo_verificar)
            if arquivo_verificar
            else None
        ),
    }



def preparar_arquivo_ordens(
    arquivo,
) -> tuple[list[dict], list[str], list[str]]:
    arquivo.seek(0)

    df_pedidos = pd.read_excel(
        arquivo,
        dtype=str,
    ).fillna("")

    arquivo.seek(0)

    erros = validar_planilha_pedidos(
        df_pedidos
    )

    if erros:
        return [], [], erros

    try:
        execucoes, alertas = preparar_execucoes(
            df_pedidos
        )
    except Exception as erro:
        return [], [], [str(erro)]

    return execucoes, alertas, []
