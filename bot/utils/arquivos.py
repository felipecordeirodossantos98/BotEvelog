from __future__ import annotations

import unicodedata
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


def normalizar_cabecalho(valor: object) -> str:
    texto = unicodedata.normalize(
        "NFKD",
        str(valor).strip(),
    )

    texto = "".join(
        caractere
        for caractere in texto
        if not unicodedata.combining(caractere)
    )

    return " ".join(
        texto.upper().split()
    )


def ler_cabecalho_excel(
    arquivo,
    linha_cabecalho: int,
) -> list[str]:
    arquivo.seek(0)

    try:
        df = pd.read_excel(
            arquivo,
            header=linha_cabecalho,
            nrows=0,
        )
    finally:
        arquivo.seek(0)

    return [
        normalizar_cabecalho(coluna)
        for coluna in df.columns
    ]


def identificar_fluxo_automacao_sac(
    arquivo,
) -> str | None:
    colunas_linha_1 = ler_cabecalho_excel(
        arquivo,
        linha_cabecalho=0,
    )

    if colunas_linha_1[:2] == [
        "SIGLA DO RESTAURANTE",
        "NUMERO DE ORDENS",
    ]:
        return "gerar_ordens_de_coleta"

    if (
        colunas_linha_1
        and colunas_linha_1[0] == "ORDEM/CTE"
    ):
        return "baixar_danfes"

    if colunas_linha_1[:2] == [
        "CTE",
        "PESO",
    ]:
        return "buscar_tdes"

    if (
        colunas_linha_1
        and colunas_linha_1[0] in {
            "NUMERO DO PEDIDO",
            "DATA",
        }
    ):
        return "baixar_relatorios_performance"

    colunas_linha_2 = ler_cabecalho_excel(
        arquivo,
        linha_cabecalho=1,
    )

    if (
        colunas_linha_2
        and colunas_linha_2[0] == "CODIGO"
    ):
        return "gerar_tickets_zendesk"

    return None



def salvar_planilha_reimportacao(
    df: pd.DataFrame,
    pasta: Path,
    prefixo: str,
    *,
    colunas_texto: tuple[str, ...] = (),
) -> Path | None:
    """
    Salva somente os registros com erro no mesmo formato aceito pelo fluxo.

    A planilha é propositalmente simples: cabeçalho na primeira linha e
    somente as colunas necessárias para que o arquivo possa ser importado
    novamente no Bot Evelog.
    """
    if df.empty:
        return None

    pasta = Path(pasta)
    pasta.mkdir(
        parents=True,
        exist_ok=True,
    )

    agora = datetime.now().strftime(
        "%Y-%m-%d_%H-%M-%S"
    )

    caminho = (
        pasta
        / f"{prefixo}_{agora}.xlsx"
    )

    df.to_excel(
        caminho,
        index=False,
    )

    if colunas_texto:
        wb = load_workbook(
            caminho
        )
        ws = wb.active

        cabecalhos = {
            str(celula.value): indice
            for indice, celula
            in enumerate(
                ws[1],
                start=1,
            )
        }

        for coluna in colunas_texto:
            indice_coluna = cabecalhos.get(
                coluna
            )

            if indice_coluna is None:
                continue

            for linha in range(
                2,
                ws.max_row + 1,
            ):
                celula = ws.cell(
                    linha,
                    indice_coluna,
                )

                if celula.value is not None:
                    celula.value = str(
                        celula.value
                    )

                celula.number_format = "@"

        wb.save(
            caminho
        )

    return caminho
